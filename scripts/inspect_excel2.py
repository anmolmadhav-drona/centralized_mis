#!/usr/bin/env python3
"""
STAGE 1 (part 2): Deeper inspection - table definition, distinct values,
number formats, special rows, duplicate LR check, summary structure.
"""
import openpyxl
from collections import Counter

PATH = "/home/z/my-project/upload/Updated MIS NPL SONIPAT.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=False)
wbv = openpyxl.load_workbook(PATH, data_only=True)

ws = wb["MIS"]
wsv = wbv["MIS"]

print("=" * 80)
print("EXCEL TABLES in workbook:")
for tbl_name, tbl in ws.tables.items():
    print(f"  Table: {tbl_name}, ref={tbl}")
    try:
        tbl_obj = ws.tables[tbl_name]
        print(f"    Columns: {[c.name for c in tbl_obj.tableColumns]}")
        print(f"    Style: {tbl_obj.tableStyleInfo.name if tbl_obj.tableStyleInfo else None}")
    except Exception as e:
        print(f"    (detail error: {e})")

print("\n" + "=" * 80)
print("SPECIAL ROWS with formulas (rows 57, 163 - computed values):")
for r in [1, 56, 57, 58, 162, 163, 164]:
    vals = [wsv.cell(row=r, column=c).value for c in range(1, 31)]
    print(f"  Row {r}: {vals}")

print("\n" + "=" * 80)
print("HIDDEN ROWS - what data do they contain?")
hidden_rows = [k for k, v in ws.row_dimensions.items() if v.hidden]
print(f"  Hidden row numbers ({len(hidden_rows)}): {hidden_rows}")
for r in hidden_rows[:10]:
    vals = [wsv.cell(row=r, column=c).value for c in range(1, 31)]
    nonempty = {i+1: v for i, v in enumerate(vals) if v is not None}
    print(f"  Row {r}: {nonempty}")

print("\n" + "=" * 80)
print("DISTINCT VALUES for categorical columns (from data rows 3..344):")
DISTINCT_COLS = {
    "PICKUP LOCATION": 2, "PARTY NAME": 3, "DESTINATION": 4, "LOAD TYPE FTL/PTL": 13,
    "DELIVERY STATUS": 16, "LR STATUS": 17, "DAMAGE": 18, "Remarsk": 24, "Remarks 1": 25,
    "VEHICLE TYPE": 22, "PLY": 23, "POD Status": 30, "MATERIAL DETAILS": 9,
    "TRANSPOTER NAME": 10, "Vendor Name": 28, "Route Code2": 29,
}
for name, c in DISTINCT_COLS.items():
    counter = Counter()
    for r in range(3, 345):
        v = wsv.cell(row=r, column=c).value
        if v is not None:
            counter[str(v).strip()] += 1
    print(f"  {name}: {len(counter)} distinct -> {dict(counter.most_common(25))}")

print("\n" + "=" * 80)
print("LR. NO. analysis:")
lr_vals = [wsv.cell(row=r, column=6).value for r in range(3, 345)]
lr_nonempty = [v for v in lr_vals if v is not None]
dup = [k for k, cnt in Counter(lr_nonempty).items() if cnt > 1]
print(f"  Total non-empty LR numbers: {len(lr_nonempty)}")
print(f"  Min: {min(lr_nonempty)}, Max: {max(lr_nonempty)}")
print(f"  Duplicate LR numbers: {dup[:20]}")
print(f"  Data rows with any data: {sum(1 for r in range(3,345) if any(wsv.cell(row=r,column=c).value is not None for c in range(1,31)))}")

print("\n" + "=" * 80)
print("NUMBER FORMATS (row 3 sample):")
for c in range(1, 31):
    cell = ws.cell(row=3, column=c)
    print(f"  Col {c} ({wsv.cell(row=2, column=c).value}): format='{cell.number_format}'")

print("\n" + "=" * 80)
print("ROW 1 & 2 header formulas:")
for c in range(1, 31):
    v = ws.cell(row=1, column=c).value
    if v is not None:
        print(f"  Row1 Col {c}: {v!r}")
hdr2 = [ws.cell(row=2, column=c).value for c in range(1, 31)]
print(f"  Row 2 headers: {hdr2}")

print("\n" + "=" * 80)
print("SUMMARY SHEET - check if real pivot or values, formulas:")
ws_sum = wb["Summary"]
formula_count = 0
for row in ws_sum.iter_rows():
    for cell in row:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            formula_count += 1
            if formula_count <= 15:
                print(f"  {cell.coordinate}: {cell.value[:200]}")
print(f"  Total formulas in Summary: {formula_count}")
# Check pivot table caches
print(f"\n  Pivot tables in Summary sheet: {len(getattr(ws_sum, '_pivots', []))}")
try:
    pivots = ws_sum._pivots
    for p in pivots:
        print(f"    Pivot: {p.name}, location={p.location.ref}")
        print(f"    Row fields: {[(f.name, f.axis) for f in p.cache.cacheFields]}")
except Exception as e:
    print(f"    pivot read: {e}")

print("\n" + "=" * 80)
print("HEADER DATES / LR DATE range:")
lr_dates = [wsv.cell(row=r, column=7).value for r in range(3, 345) if wsv.cell(row=r, column=7).value]
if lr_dates:
    print(f"  LR DATE min: {min(lr_dates)}, max: {max(lr_dates)}")
inv = [wsv.cell(row=r, column=5).value for r in range(3, 345) if wsv.cell(row=r, column=5).value]
print(f"  INVOICE samples: {inv[:5]}")
ints_in_inv = [v for v in inv if isinstance(v, int)]
print(f"  INVOICE integer values: {ints_in_inv[:10]}")

print("\n" + "=" * 80)
print("FROZEN PANES / AUTOFILTER:")
print(f"  MIS freeze_panes: {ws.freeze_panes}")
print(f"  MIS auto_filter: {ws.auto_filter.ref}")
print(f"  Summary freeze_panes: {ws_sum.freeze_panes}")
