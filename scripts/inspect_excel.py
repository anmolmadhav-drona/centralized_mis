#!/usr/bin/env python3
"""
STAGE 1: Reverse-engineer 'Updated MIS NPL SONIPAT.xlsx'
Inspect: sheets, dimensions, columns, types, formulas, validation,
merged cells, formatting, hidden rows/cols, summary logic.
"""
import openpyxl
import json
from datetime import datetime, date

PATH = "/home/z/my-project/upload/Updated MIS NPL SONIPAT.xlsx"

def fmt_val(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return str(v)
    if isinstance(v, float):
        if v == int(v):
            return int(v)
        return v
    return v

print("=" * 80)
print("WORKBOOK INSPECTION:", PATH)
print("=" * 80)

# Load with formulas (read-only off to access formulas & validation)
wb = openpyxl.load_workbook(PATH, data_only=False)
wb_data = openpyxl.load_workbook(PATH, data_only=True)

print("\nSHEET NAMES:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    wsd = wb_data[sheet_name]
    print("\n" + "#" * 80)
    print(f"SHEET: '{sheet_name}'")
    print(f"  Dimensions: {ws.dimensions}")
    print(f"  Max row: {ws.max_row}, Max col: {ws.max_column}")
    print(f"  Sheet state: {ws.sheet_state}")

    # Hidden rows / columns
    hidden_cols = [k for k, v in ws.column_dimensions.items() if v.hidden]
    hidden_rows = [k for k, v in ws.row_dimensions.items() if v.hidden]
    if hidden_cols:
        print(f"  Hidden columns: {hidden_cols}")
    if hidden_rows:
        print(f"  Hidden rows (first 20): {hidden_rows[:20]}")

    # Merged cells
    if ws.merged_cells.ranges:
        print(f"  Merged cells: {[str(r) for r in ws.merged_cells.ranges][:30]}")

    # Data validation (dropdowns)
    try:
        dvs = ws.data_validations.dataValidation
        for dv in dvs:
            print(f"  DATA VALIDATION: type={dv.type}, formula1={dv.formula1}, ranges={dv.sqref}")
    except Exception as e:
        print(f"  (data validation read error: {e})")

    # Conditional formatting
    try:
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            for rule in rules:
                print(f"  CONDITIONAL FORMAT: range={rng.sqref}, type={rule.type}, formula={getattr(rule, 'formula', None)}")
    except Exception as e:
        print(f"  (conditional format read error: {e})")

    # Scan for formulas (sample)
    formulas = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200)):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                col_letter = cell.column_letter
                formulas.setdefault(col_letter, []).append(
                    {"row": cell.row, "formula": cell.value}
                )
    if formulas:
        print("  FORMULAS by column:")
        for col, items in formulas.items():
            print(f"    Column {col}: {len(items)} formulas. Examples:")
            for it in items[:3]:
                print(f"      row {it['row']}: {it['formula'][:160]}")

    # Print first rows (headers + sample data)
    print("\n  FIRST 8 ROWS (computed values):")
    for r in range(1, min(ws.max_row, 8) + 1):
        row_vals = []
        for c in range(1, ws.max_column + 1):
            v = wsd.cell(row=r, column=c).value
            row_vals.append(fmt_val(v))
        print(f"    Row {r}: {row_vals}")

    # Data type profiling per column (skip header row(s))
    print("\n  COLUMN PROFILE (rows 2..min(300,max_row)):")
    max_r = min(ws.max_row, 300)
    for c in range(1, ws.max_column + 1):
        header = wsd.cell(row=1, column=c).value
        type_counts = {}
        samples = []
        for r in range(2, max_r + 1):
            v = wsd.cell(row=r, column=c).value
            if v is None:
                type_counts["empty"] = type_counts.get("empty", 0) + 1
                continue
            t = type(v).__name__
            type_counts[t] = type_counts.get(t, 0) + 1
            if len(samples) < 5:
                samples.append(fmt_val(v))
        print(f"    Col {c} [{header!r}]: {type_counts} | samples: {samples}")

    # Last rows (often totals / summary)
    print("\n  LAST 5 ROWS (computed values):")
    if ws.max_row > 8:
        for r in range(max(1, ws.max_row - 4), ws.max_row + 1):
            row_vals = []
            for c in range(1, ws.max_column + 1):
                v = wsd.cell(row=r, column=c).value
                row_vals.append(fmt_val(v))
            print(f"    Row {r}: {row_vals}")

print("\n" + "=" * 80)
print("INSPECTION COMPLETE")
